"""Compare real XLS/XLSX inputs and preserve reader/layout boundaries."""
from decimal import Decimal
import json
from pathlib import Path

from openpyxl import Workbook
import pandas
import pytest
import python_calamine

from app.parsers import parse_excel, UnsupportedLayoutError

ROOT = Path(__file__).parent / 'fixtures' / 'parser_equivalence'


def decimal_text(value):
    if isinstance(value, Decimal):
        return str(value)
    raise TypeError(type(value).__name__)


def check_reference(case, extension):
    result = parse_excel(ROOT / case / f'input.{extension}')
    assert json.loads(json.dumps(result, default=decimal_text)) == json.loads(
        (ROOT / case / 'expected.json').read_text()
    )
    for person in result['colaboradores']:
        assert isinstance(person['liquido'], Decimal)
        assert all(isinstance(r['valor'], Decimal) for r in person['rubricas'])


@pytest.mark.parametrize('case,extension', [
    ('extrato_xls', 'xls'), ('extrato_xlsx', 'xlsx'),
    ('extrato_xlsx_edge', 'xlsx'), ('extrato_xlsx_empty', 'xlsx'),
])
def test_excel_matches_frozen_reference(case, extension):
    check_reference(case, extension)


def fail_primary(*args, **kwargs):
    raise OSError('primary reader failed')


@pytest.mark.parametrize('extension', ['xls', 'xlsx'])
def test_excel_fallback_matches_reference(monkeypatch, extension):
    monkeypatch.setattr(python_calamine, 'load_workbook', fail_primary)
    check_reference(f'extrato_{extension}', extension)


def test_excel_propagates_failure_when_both_readers_fail(monkeypatch):
    def fail_fallback(*args, **kwargs):
        raise ValueError('fallback reader failed')

    monkeypatch.setattr(python_calamine, 'load_workbook', fail_primary)
    monkeypatch.setattr(pandas, 'read_excel', fail_fallback)
    with pytest.raises(ValueError, match='fallback reader failed'):
        parse_excel(ROOT / 'extrato_xlsx' / 'input.xlsx')


def write_workbook(path, first_rows, second_rows=()):
    workbook = Workbook()
    for row in first_rows:
        workbook.active.append(row)
    if second_rows:
        sheet = workbook.create_sheet('Outra empresa')
        for row in second_rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize('fallback', [False, True])
def test_excel_uses_first_sheet_only(tmp_path, monkeypatch, fallback):
    if fallback:
        monkeypatch.setattr(python_calamine, 'load_workbook', fail_primary)
    path = tmp_path / 'two_sheets.xlsx'
    write_workbook(path, [['EXTRATO MENSAL'], ['Competência:', '09/2026']], [
        ['EXTRATO MENSAL'], ['Competência:', '10/2026'],
        ['Empr.:', '0002', 'OUTRA PESSOA', 'CPF:', '111.111.111-11'],
        ['001', 'SALARIO BASE', '30,00', '500,00', 'P'],
    ])
    result = parse_excel(path)
    assert result['competencia'] == '09/2026'
    assert result['colaboradores'] == []


def test_extrato_title_takes_precedence_over_thirteenth_salary(tmp_path):
    path = tmp_path / 'monthly.xlsx'
    write_workbook(path, [
        ['EXTRATO MENSAL', 'PROVISÃO DE 13º SALÁRIO'],
        ['Competência:', '09/2026'],
        ['Empr.:', '0001', 'ANA SINTÉTICA', 'CPF:', '000.000.000-00'],
        ['001', 'SALARIO BASE', '30,00', '500,00', 'P'],
    ])
    result = parse_excel(path)
    assert result['tipo_importacao'] == 'Extrato'
    assert result['colaboradores'][0]['rubricas'][0]['valor'] == Decimal('500.00')


@pytest.mark.parametrize('title', [
    'PROVISÃO DE 13º SALÁRIO', 'PROVISAO DE 13 SALARIO',
    'PROVISÃO DE DÉCIMO TERCEIRO', 'PROVISAO DE DECIMO TERCEIRO',
    'PROVISÃO DE FÉRIAS', 'PROVISAO DE FERIAS', 'DETALHAMENTO DA PROVISÃO DE FÉRIAS',
])
def test_provision_layouts_are_rejected(tmp_path, title):
    path = tmp_path / 'provision.xlsx'
    write_workbook(path, [[title]])
    with pytest.raises(UnsupportedLayoutError):
        parse_excel(path)
